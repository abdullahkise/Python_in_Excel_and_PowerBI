"""
BUS_Matrix.xlsx üretici (openpyxl).

Amaç : O2'de Excel'de canlı çizilecek BUS Matrix'in HAZIR kopyasını (Plan B / referans) üretmek.
       Sayfalar: BUS_Matrix (iş süreci x ortak boyut, formüllü sayaçlar), Grain_ve_FactSales, Boyutlar.
Bağımlılık: openpyxl
Kullanım  : python uret_bus_matrix_xlsx.py [çıktı yolu]   (varsayılan: ..\\BUS_Matrix.xlsx)
"""

import sys
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

cikti = Path(sys.argv[1]) if len(sys.argv) > 1 else Path(__file__).resolve().parents[1] / "BUS_Matrix.xlsx"

KOYU = PatternFill("solid", fgColor="1F3864")
MAVI = PatternFill("solid", fgColor="2F75B5")
YESIL = PatternFill("solid", fgColor="70AD47")
GRI = PatternFill("solid", fgColor="D9D9D9")
ACIK = PatternFill("solid", fgColor="DDEBF7")
SARI = PatternFill("solid", fgColor="FFF2CC")
BEYAZ_KALIN = Font(bold=True, color="FFFFFF", size=11)
KALIN = Font(bold=True)
ince = Side(style="thin", color="999999")
CERCEVE = Border(left=ince, right=ince, top=ince, bottom=ince)
ORTA = Alignment(horizontal="center", vertical="center", wrap_text=True)
SOL = Alignment(horizontal="left", vertical="center", wrap_text=True)


def baslik_satiri(ws, satir, degerler, dolgu=KOYU):
    for i, v in enumerate(degerler, start=1):
        c = ws.cell(row=satir, column=i, value=v)
        c.fill, c.font, c.alignment, c.border = dolgu, BEYAZ_KALIN, ORTA, CERCEVE


wb = Workbook()

# ---------------------------------------------------------------- 1) BUS Matrix
ws = wb.active
ws.title = "BUS_Matrix"
ws["A1"] = "BUS Matrix – Pedalova Bisiklet A.Ş."
ws["A1"].font = Font(bold=True, size=16, color="1F3864")
ws["A2"] = "Satırlar = iş süreçleri, sütunlar = ortak (conformed) boyutlar. X = süreç o boyutla analiz edilir. FactSales yalnızca 'Satış' satırından türetilir."
ws["A2"].font = Font(italic=True, color="595959")
ws.merge_cells("A2:L2")

baslik = ["İş Süreci", "Fakt tablosu / Grain", "Tarih", "Müşteri", "Ürün", "Mağaza", "Kanal", "Ödeme", "Kampanya",
          "Boyut sayısı", "Durum", "Not"]
baslik_satiri(ws, 4, baslik)
surecler = [
    ("Satış (sipariş satırı)", "FactSales – bir sipariş satırı", "X", "X", "X", "X", "X", "X", "X", "Bu projede", "FactSales buradan doğar"),
    ("İade", "FactReturns – bir iade satırı", "X", "X", "X", "X", "X", "X", "", "Gelecek", "Mevcut boyutlar aynen kullanılır"),
    ("Envanter hareketi", "FactInventory – ürün x mağaza x gün", "X", "", "X", "X", "", "", "", "Gelecek", "Müşteri/ödeme yok"),
    ("Kampanya bütçesi", "FactCampaignBudget – kampanya x ay", "X", "", "X", "", "X", "", "X", "Gelecek", "Mağaza/müşteri yok"),
]
for r, s in enumerate(surecler, start=5):
    ws.cell(row=r, column=1, value=s[0]).font = KALIN
    ws.cell(row=r, column=2, value=s[1])
    for j in range(7):
        c = ws.cell(row=r, column=3 + j, value=s[2 + j] or None)
        c.alignment = ORTA
        if s[2 + j] == "X":
            c.fill = YESIL if s[9] == "Bu projede" else MAVI
            c.font = BEYAZ_KALIN
    ws.cell(row=r, column=10, value=f'=COUNTIF(C{r}:I{r},"X")').alignment = ORTA
    ws.cell(row=r, column=11, value=s[9]).alignment = ORTA
    ws.cell(row=r, column=12, value=s[10])
    for col in range(1, 13):
        ws.cell(row=r, column=col).border = CERCEVE
        if col in (1, 2, 11, 12):
            ws.cell(row=r, column=col).alignment = SOL
    if s[9] == "Bu projede":
        for col in (1, 2, 10, 11, 12):
            ws.cell(row=r, column=col).fill = SARI

ws.cell(row=9, column=1, value="Boyutu kullanan süreç sayısı").font = KALIN
for j in range(7):
    L = get_column_letter(3 + j)
    c = ws.cell(row=9, column=3 + j, value=f'=COUNTIF({L}5:{L}8,"X")')
    c.alignment, c.fill, c.font = ORTA, GRI, KALIN
ws.cell(row=9, column=1).fill = GRI
ws.merge_cells("A9:B9")
for col in range(1, 10):
    ws.cell(row=9, column=col).border = CERCEVE

ws["A11"] = "Nasıl okunur?"
ws["A11"].font = Font(bold=True, size=12, color="1F3864")
okuma = [
    "1) Önce SATIR: hangi iş sürecini ölçüyoruz? (Satış)  2) Sonra GRAIN: fakt tablosunda bir satır neyi temsil eder? (bir sipariş satırı)",
    "3) Sonra SÜTUN: o süreci hangi bakış açılarıyla (boyutlarla) inceleyeceğiz?  Yeşil X'ler FactSales'in yabancı anahtarlarıdır.",
    "4) Altındaki sayılar: bir boyutu kaç süreç kullanıyor? 2+ ise 'ortak boyut' (conformed dimension) – tek tanım, her yerde aynı anlam.",
    "5) Mavi X'ler yarın eklenecek süreçlerdir: Tarih/Ürün/Mağaza boyutları YENİDEN yazılmaz, aynen kullanılır.",
]
for i, t in enumerate(okuma, start=12):
    ws.cell(row=i, column=1, value=t)
    ws.merge_cells(start_row=i, start_column=1, end_row=i, end_column=12)
    ws.cell(row=i, column=1).alignment = SOL

ws["A17"] = "Kimball vs Inmon (genel kültür)"
ws["A17"].font = Font(bold=True, size=12, color="1F3864")
ws["A18"] = ("Kimball (aşağıdan yukarı): süreç odaklı boyutsal modeller (star schema) -> BUS Matrix ile ortak boyutlar -> kurumsal DW. "
             "Inmon (yukarıdan aşağı): önce normalize kurumsal DW (3NF), sonra department data mart'lar. Bu proje: Kimball + Medallion.")
ws.merge_cells("A18:L19")
ws["A18"].alignment = SOL

for col, w in zip("ABCDEFGHIJKL", [26, 38, 9, 10, 9, 10, 9, 9, 11, 12, 12, 36]):
    ws.column_dimensions[col].width = w
ws.row_dimensions[4].height = 28
ws.freeze_panes = "C5"

# ---------------------------------------------------------------- 2) Grain ve FactSales
g = wb.create_sheet("Grain_ve_FactSales")
g["A1"] = "Grain ve FactSales"
g["A1"].font = Font(bold=True, size=16, color="1F3864")
g["A2"] = "GRAIN: gold.FactSales'te bir satır = bir siparişin bir ürün satırı (OrderID x OrderLineID)."
g["A2"].font = Font(bold=True, color="C00000")
g.merge_cells("A2:F2")
g["A3"] = "Neden 'sipariş' değil 'sipariş satırı'? Ürün, adet, fiyat ve indirim SATIR düzeyindedir; sipariş düzeyine toplamak bilgi kaybettirir (tersi mümkündür)."
g.merge_cells("A3:F3")
baslik_satiri(g, 5, ["Fakt kolonu", "Tür", "Toplanabilirlik", "Açıklama", "Silver kaynağı", "DAX / KPI notu"])
fakt = [
    ("DateKey", "Boyut anahtarı", "-", "DimDate'e FK (yyyymmdd)", "SiparisTarihi", ""),
    ("CustomerKey", "Boyut anahtarı", "-", "DimCustomer'a FK; -1 = Bilinmiyor", "MusteriID", ""),
    ("ProductKey", "Boyut anahtarı", "-", "DimProduct'a FK", "UrunID", ""),
    ("StoreKey", "Boyut anahtarı", "-", "DimStore'a FK", "MagazaID", ""),
    ("ChannelKey", "Boyut anahtarı", "-", "DimChannel'a FK", "SatisKanali", ""),
    ("PaymentKey", "Boyut anahtarı", "-", "DimPayment'a FK", "OdemeTipi", ""),
    ("CampaignKey", "Boyut anahtarı", "-", "DimCampaign'a FK", "KampanyaAdi", ""),
    ("OrderID", "Degenerate boyut", "-", "Sipariş no (kendi boyut tablosu yok); grain'in parçası", "SiparisID", "Sipariş sayısı = DISTINCTCOUNT(OrderID)"),
    ("OrderLineID", "Degenerate boyut", "-", "Siparişteki satır sırası", "SiparisSatirID", ""),
    ("Quantity", "Ölçü", "Toplanabilir (additive)", "Satırdaki ürün adedi (sipariş sayısı DEĞİL)", "SiparisAdedi", "Satış Adedi = SUM"),
    ("UnitPrice", "Ölçü", "Toplanamaz (non-additive)", "Liste birim fiyatı", "BirimFiyat", "Toplanmaz; ortalama veya adetle ağırlıklı kullan"),
    ("UnitCost", "Ölçü", "Toplanamaz (non-additive)", "Birim maliyet", "UrunMaliyeti", ""),
    ("DiscountRate", "Ölçü", "Toplanamaz (non-additive)", "Satır indirim oranı", "IndirimOrani", "Ortalama indirim % = İndirim Tutarı / Brüt Satış (ağırlıklı)"),
    ("GrossAmount", "Ölçü", "Toplanabilir", "Quantity x UnitPrice", "hesaplanır", "Brüt Satış = SUM"),
    ("DiscountAmount", "Ölçü", "Toplanabilir", "GrossAmount x DiscountRate", "hesaplanır", "İndirim Tutarı = SUM"),
    ("NetAmount", "Ölçü", "Toplanabilir", "GrossAmount - DiscountAmount", "hesaplanır", "Net Satış = SUM"),
    ("CostAmount", "Ölçü", "Toplanabilir", "Quantity x UnitCost", "hesaplanır", "Maliyet = SUM"),
    ("GrossProfit", "Ölçü", "Toplanabilir", "NetAmount - CostAmount", "hesaplanır", "Brüt Kâr = SUM; Marj % = Kâr / Net (oran: toplanmaz)"),
    ("CustomerAgeAtOrder", "Olay anı özelliği", "Toplanamaz", "Müşterinin o siparişteki yaşı (fakt'ta durur, boyutta değil)", "MusteriYas", "Regresyonda X değişkeni"),
    ("SourceFileName, BatchID, GoldLoadDateTime", "Lineage", "-", "Veri yaşam döngüsü izi", "SourceFileName, BatchID", ""),
]
for r, satir in enumerate(fakt, start=6):
    for c, v in enumerate(satir, start=1):
        cell = g.cell(row=r, column=c, value=v)
        cell.border, cell.alignment = CERCEVE, SOL
    tur = satir[1]
    g.cell(row=r, column=2).fill = ACIK if tur == "Boyut anahtarı" else YESIL if tur == "Ölçü" else SARI if "Degenerate" in tur else GRI
for col, w in zip("ABCDEF", [30, 18, 26, 52, 22, 48]):
    g.column_dimensions[col].width = w
g.freeze_panes = "A6"

# ---------------------------------------------------------------- 3) Boyutlar
b = wb.create_sheet("Boyutlar")
b["A1"] = "Boyutlar (Dimensions)"
b["A1"].font = Font(bold=True, size=16, color="1F3864")
baslik_satiri(b, 3, ["Boyut", "Yapay anahtar", "İş anahtarı", "Öznitelikler", "Kaynak", "Notlar"])
boyutlar = [
    ("DimDate", "DateKey (yyyymmdd)", "FullDate", "Year, Quarter, QuarterName, MonthNo, MonthName, MonthShort, YearMonth, YearMonthKey, WeekOfYear, DayOfMonth, DayOfWeekNo, DayName, IsWeekend",
     "Üretim scripti (2022-01-01..2026-12-31)", "Tam yıllar, kesintisiz: Time Intelligence şartı"),
    ("DimCustomer", "CustomerKey (-1 = Bilinmiyor)", "CustomerID", "CustomerName, Gender, City, Region, MembershipYear, LatestAge, AgeGroup",
     "silver.Sales (son sipariş)", "SCD Tip 1; segment burada YOK (Excel K-Means çıktısı ayrı)"),
    ("DimProduct", "ProductKey", "ProductID", "ProductName, Category, SubCategory, Brand", "silver.RefProduct", "Ana veriden: satılmayan ürünler de var"),
    ("DimStore", "StoreKey", "StoreID", "StoreName, City, Region, StoreType", "silver.RefStore", "S900 = Dijital Kanal Deposu (web/mobil)"),
    ("DimChannel", "ChannelKey", "ChannelName", "ChannelGroup (Fiziksel/Dijital)", "silver.Sales", "Küçük gruplama boyutu"),
    ("DimPayment", "PaymentKey", "PaymentType", "PaymentGroup (Kart/Nakit/Banka Transferi)", "silver.Sales", "Küçük gruplama boyutu"),
    ("DimCampaign", "CampaignKey", "CampaignName", "CampaignType, CampaignSeason", "silver.RefCampaign", "Ramazan vs Kurban Bayramı kıyası için"),
]
for r, satir in enumerate(boyutlar, start=4):
    for c, v in enumerate(satir, start=1):
        cell = b.cell(row=r, column=c, value=v)
        cell.border, cell.alignment = CERCEVE, SOL
    b.cell(row=r, column=1).font = KALIN
for col, w in zip("ABCDEF", [16, 28, 16, 70, 34, 48]):
    b.column_dimensions[col].width = w

cikti.parent.mkdir(parents=True, exist_ok=True)
wb.save(cikti)
print("Yazıldı:", cikti)
